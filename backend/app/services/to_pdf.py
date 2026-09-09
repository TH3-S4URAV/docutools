import os
import sys
import shutil
import subprocess
from pathlib import Path
from typing import Union, List
from PIL import Image
import pymupdf
import docx
from pptx import Presentation
import openpyxl

def _try_libreoffice(input_path: str, output_dir: str) -> bool:
    """Attempt conversion using headless LibreOffice if installed (cross-platform / Linux / Docker)."""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice and sys.platform == "win32":
        lo_paths = [
            r"C:\Program Files\LibreOffice\program\soffice.exe",
            r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"
        ]
        for p in lo_paths:
            if os.path.exists(p):
                soffice = p
                break
    if soffice:
        try:
            cmd = [soffice, "--headless", "--convert-to", "pdf", "--outdir", output_dir, input_path]
            res = subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=30)
            return res.returncode == 0
        except Exception:
            return False
    return False

def docx_to_pdf(file_path: Union[str, Path], output_path: Union[str, Path]) -> str:
    """Convert DOCX to PDF using LibreOffice, Windows Word COM, or pure Python fallback."""
    abs_input = str(Path(file_path).resolve())
    abs_output = str(Path(output_path).resolve())
    out_dir = str(Path(output_path).parent.resolve())

    # 1. Try LibreOffice
    if _try_libreoffice(abs_input, out_dir):
        expected = Path(out_dir) / f"{Path(abs_input).stem}.pdf"
        if expected.exists() and expected != Path(abs_output):
            expected.rename(abs_output)
        if os.path.exists(abs_output) and os.path.getsize(abs_output) > 0:
            return abs_output

    # 2. Try Windows Word COM
    if sys.platform == "win32":
        try:
            import win32com.client
            import pythoncom
            pythoncom.CoInitialize()
            word = None
            doc = None
            try:
                word = win32com.client.DispatchEx("Word.Application")
                word.Visible = False
                word.DisplayAlerts = False
                doc = word.Documents.Open(abs_input)
                doc.SaveAs2(abs_output, FileFormat=17)
                doc.Close(False)
                word.Quit()
                del doc
                del word
                if os.path.exists(abs_output) and os.path.getsize(abs_output) > 0:
                    return abs_output
            finally:
                pythoncom.CoUninitialize()
        except Exception:
            pass

    # 3. Pure Python fallback
    doc_file = docx.Document(abs_input)
    pdf_doc = pymupdf.open()
    page = pdf_doc.new_page()
    p_cursor = 50

    for para in doc_file.paragraphs:
        text = para.text.strip()
        if not text:
            p_cursor += 15
            continue
        if p_cursor > page.rect.height - 50:
            page = pdf_doc.new_page()
            p_cursor = 50
        page.insert_text(pymupdf.Point(50, p_cursor), text, fontsize=11, fontname="helv")
        p_cursor += 18

    pdf_doc.save(abs_output, garbage=3, deflate=True)
    pdf_doc.close()
    return abs_output

def pptx_to_pdf(file_path: Union[str, Path], output_path: Union[str, Path]) -> str:
    """Convert PPTX to PDF using LibreOffice, Windows PowerPoint COM, or pure Python fallback."""
    abs_input = str(Path(file_path).resolve())
    abs_output = str(Path(output_path).resolve())
    out_dir = str(Path(output_path).parent.resolve())

    # 1. Try LibreOffice
    if _try_libreoffice(abs_input, out_dir):
        expected = Path(out_dir) / f"{Path(abs_input).stem}.pdf"
        if expected.exists() and expected != Path(abs_output):
            expected.rename(abs_output)
        if os.path.exists(abs_output) and os.path.getsize(abs_output) > 0:
            return abs_output

    # 2. Try Windows PowerPoint COM
    if sys.platform == "win32":
        try:
            import win32com.client
            import pythoncom
            pythoncom.CoInitialize()
            ppt = None
            prs = None
            try:
                ppt = win32com.client.DispatchEx("PowerPoint.Application")
                prs = ppt.Presentations.Open(abs_input, WithWindow=False)
                prs.SaveAs(abs_output, 32)
                prs.Close()
                ppt.Quit()
                del prs
                del ppt
                if os.path.exists(abs_output) and os.path.getsize(abs_output) > 0:
                    return abs_output
            finally:
                pythoncom.CoUninitialize()
        except Exception:
            pass

    # 3. Fallback
    prs_file = Presentation(abs_input)
    pdf_doc = pymupdf.open()
    width = prs_file.slide_width.pt
    height = prs_file.slide_height.pt

    for slide in prs_file.slides:
        page = pdf_doc.new_page(width=width, height=height)
        y = 50
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    txt = paragraph.text.strip()
                    if txt:
                        page.insert_text(pymupdf.Point(50, y), txt, fontsize=12)
                        y += 20

    pdf_doc.save(abs_output, garbage=3, deflate=True)
    pdf_doc.close()
    return abs_output

def xlsx_to_pdf(file_path: Union[str, Path], output_path: Union[str, Path]) -> str:
    """Convert XLSX to PDF using LibreOffice, Windows Excel COM, or pure Python fallback."""
    abs_input = str(Path(file_path).resolve())
    abs_output = str(Path(output_path).resolve())
    out_dir = str(Path(output_path).parent.resolve())

    # 1. Try LibreOffice
    if _try_libreoffice(abs_input, out_dir):
        expected = Path(out_dir) / f"{Path(abs_input).stem}.pdf"
        if expected.exists() and expected != Path(abs_output):
            expected.rename(abs_output)
        if os.path.exists(abs_output) and os.path.getsize(abs_output) > 0:
            return abs_output

    # 2. Try Windows Excel COM
    if sys.platform == "win32":
        try:
            import win32com.client
            import pythoncom
            pythoncom.CoInitialize()
            excel = None
            wb = None
            try:
                excel = win32com.client.DispatchEx("Excel.Application")
                excel.Visible = False
                excel.DisplayAlerts = False
                wb = excel.Workbooks.Open(abs_input)
                wb.ExportAsFixedFormat(0, abs_output)
                wb.Close(False)
                excel.Quit()
                del wb
                del excel
                if os.path.exists(abs_output) and os.path.getsize(abs_output) > 0:
                    return abs_output
            finally:
                pythoncom.CoUninitialize()
        except Exception:
            pass

    # 3. Fallback: openpyxl to PDF table rendering
    wb_file = openpyxl.load_workbook(abs_input, data_only=True)
    pdf_doc = pymupdf.open()

    for sheet in wb_file.worksheets:
        page = pdf_doc.new_page(width=842, height=595)
        y = 40
        page.insert_text(pymupdf.Point(40, y), f"Sheet: {sheet.title}", fontsize=14)
        y += 25
        for row in sheet.iter_rows(values_only=True):
            if any(row):
                row_str = " | ".join(str(cell) if cell is not None else "" for cell in row[:10])
                if y > 560:
                    page = pdf_doc.new_page(width=842, height=595)
                    y = 40
                page.insert_text(pymupdf.Point(40, y), row_str, fontsize=9)
                y += 16

    pdf_doc.save(abs_output, garbage=3, deflate=True)
    pdf_doc.close()
    return abs_output

def images_to_pdf(image_paths: List[Union[str, Path]], output_path: Union[str, Path],
                  orientation: str = "portrait", margin: float = 20.0) -> str:
    if not image_paths:
        raise ValueError("At least one image is required.")

    pdf_doc = pymupdf.open()
    try:
        for img_path in image_paths:
            img = Image.open(str(img_path))
            img_w, img_h = img.size
            img.close()

            page_w = img_w + 2 * margin
            page_h = img_h + 2 * margin

            page = pdf_doc.new_page(width=page_w, height=page_h)
            rect = pymupdf.Rect(margin, margin, page_w - margin, page_h - margin)
            page.insert_image(rect, filename=str(img_path))

        pdf_doc.save(str(output_path), garbage=3, deflate=True)
        return str(output_path)
    finally:
        pdf_doc.close()

def html_to_pdf(html_source: Union[str, Path], output_path: Union[str, Path]) -> str:
    abs_output = str(output_path)
    temp_file = None
    if isinstance(html_source, str) and ("<html" in html_source.lower() or "<div" in html_source.lower() or "<!doctype" in html_source.lower()):
        temp_file = Path(output_path).parent / f"temp_{os.getpid()}_{hash(html_source)}.html"
        with open(temp_file, "w", encoding="utf-8") as f:
            f.write(html_source)
        source_path = str(temp_file)
    else:
        source_path = str(html_source)

    try:
        doc = pymupdf.open(source_path)
        pdf_bytes = doc.convert_to_pdf()
        pdf_doc = pymupdf.open("pdf", pdf_bytes)
        pdf_doc.save(abs_output, garbage=3, deflate=True)
        pdf_doc.close()
        doc.close()
        return abs_output
    finally:
        if temp_file and temp_file.exists():
            temp_file.unlink(missing_ok=True)
