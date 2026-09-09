import sys
import os
from pathlib import Path
import pytest
import pymupdf
from PIL import Image, ImageDraw
import docx
from pptx import Presentation
from pptx.util import Inches
import openpyxl

# Add backend directory to sys.path
backend_dir = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(backend_dir))

@pytest.fixture(scope="session")
def test_dir(tmp_path_factory):
    d = tmp_path_factory.mktemp("docutools_test_files")
    return d

@pytest.fixture(scope="session")
def sample_pdf(test_dir):
    """Generate a valid 3-page PDF with text and an embedded image."""
    pdf_path = test_dir / "sample.pdf"
    doc = pymupdf.open()
    
    # Page 1: Title and paragraph
    p1 = doc.new_page(width=595, height=842)
    p1.insert_text(pymupdf.Point(50, 80), "DocuTools Test Document - Page 1", fontsize=18)
    p1.insert_text(pymupdf.Point(50, 120), "This is a test paragraph for automated document verification.", fontsize=11)
    
    # Draw a colored rectangle and insert a small test image
    img = Image.new("RGB", (200, 100), color=(73, 109, 137))
    d = ImageDraw.Draw(img)
    d.text((20, 40), "Test Image", fill=(255, 255, 0))
    img_path = test_dir / "test_embedded.png"
    img.save(str(img_path))
    p1.insert_image(pymupdf.Rect(50, 150, 250, 250), filename=str(img_path))

    # Page 2: Table-like content
    p2 = doc.new_page(width=595, height=842)
    p2.insert_text(pymupdf.Point(50, 80), "Page 2 - Structured Table", fontsize=16)
    p2.insert_text(pymupdf.Point(50, 120), "ID | Name | Department | Score", fontsize=11)
    p2.insert_text(pymupdf.Point(50, 140), "1  | Alice| Engineering| 95", fontsize=11)
    p2.insert_text(pymupdf.Point(50, 160), "2  | Bob  | Design     | 88", fontsize=11)

    # Page 3: Summary text
    p3 = doc.new_page(width=595, height=842)
    p3.insert_text(pymupdf.Point(50, 80), "Page 3 - Conclusion and Notes", fontsize=16)
    p3.insert_text(pymupdf.Point(50, 120), "Final page for multi-page split and extraction testing.", fontsize=11)

    doc.save(str(pdf_path))
    doc.close()
    return pdf_path

@pytest.fixture(scope="session")
def sample_docx(test_dir):
    docx_path = test_dir / "sample.docx"
    doc = docx.Document()
    doc.add_heading("DocuTools Sample Word File", 0)
    doc.add_paragraph("This is a paragraph inside sample docx for conversion testing.")
    doc.save(str(docx_path))
    return docx_path

@pytest.fixture(scope="session")
def sample_pptx(test_dir):
    pptx_path = test_dir / "sample.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Sample Presentation"
    slide.placeholders[1].text = "Subtitle for PPTX conversion testing"
    prs.save(str(pptx_path))
    return pptx_path

@pytest.fixture(scope="session")
def sample_xlsx(test_dir):
    xlsx_path = test_dir / "sample.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Item", "Quantity", "Price"])
    ws.append(["Widget A", 10, 19.99])
    ws.append(["Widget B", 25, 4.50])
    wb.save(str(xlsx_path))
    return xlsx_path

@pytest.fixture(scope="session")
def sample_images(test_dir):
    img_paths = []
    for i in range(2):
        path = test_dir / f"image_{i+1}.png"
        img = Image.new("RGB", (300, 200), color=(100 + i * 40, 150, 200))
        d = ImageDraw.Draw(img)
        d.text((40, 90), f"Sample Image {i+1}", fill=(255, 255, 255))
        img.save(str(path))
        img_paths.append(path)
    return img_paths
